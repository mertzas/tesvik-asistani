"""Başvuru dosyası Word çıktısı (2026-10-08): kontrol listesi + başvuru dönemleri + ön taslak tek .docx dosyasında.

Kurum formları çoğunlukla Word; İKAS paneli iframe'inde yazdırma güvenilir değil. Girdi, kontrol listesi uç noktasının
yanıtıdır (app/basvuru_listesi._yanit); yeni veri üretilmez. Taslak Markdown'ının küçük bir alt kümesi Word'e çevrilir:
#/## başlık, "> " alıntı, "- [ ]"/"- [x]" onay kutusu, "- " madde, "| a | b |" tablo, **kalın**.

    python -m app.basvuru_docx --self-test
"""
from __future__ import annotations

import io
import re
import sys
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

KUTU_BOS, KUTU_DOLU = "☐", "☑"
TUR_BASLIK = {"sart": "Şartlar", "belge": "Hazırlanacak belgeler", "basvuru": "Başvuru adımı"}
UYARI = ("Bu dosya Teşvik Asistanı ile hazırlanmış bir çalışma belgesidir; resmi başvuru formu değildir. Bilgileri "
         "kurumun güncel duyurusu ve formuyla karşılaştırın.")
_KALIN = re.compile(r"\*\*(.+?)\*\*")


def _satir_ici(paragraf, metin: str) -> None:
    """**kalın** işaretlerini koruyarak metni paragrafa ekler."""
    konum = 0
    for m in _KALIN.finditer(metin):
        if m.start() > konum:
            paragraf.add_run(metin[konum:m.start()])
        paragraf.add_run(m.group(1)).bold = True
        konum = m.end()
    if konum < len(metin):
        paragraf.add_run(metin[konum:])


def _tablo_hucreleri(satir: str) -> list[str]:
    return [h.strip() for h in satir.strip().strip("|").split("|")]


def markdown_ekle(doc, metin: str) -> None:
    satirlar = (metin or "").splitlines()
    i = 0
    while i < len(satirlar):
        s = satirlar[i].rstrip()
        if not s.strip():
            i += 1
            continue
        if s.lstrip().startswith("|"):
            blok = []
            while i < len(satirlar) and satirlar[i].lstrip().startswith("|"):
                if not re.fullmatch(r"\s*\|?[\s:|-]+\|?\s*", satirlar[i]):  # ayırıcı satırı (|---|---|) atla
                    blok.append(_tablo_hucreleri(satirlar[i]))
                i += 1
            sutun = max(len(r) for r in blok)
            tablo = doc.add_table(rows=len(blok), cols=sutun)
            tablo.style = "Table Grid"
            for r, hucreler in enumerate(blok):
                for c in range(sutun):
                    hucre = tablo.cell(r, c)
                    hucre.text = ""
                    _satir_ici(hucre.paragraphs[0], hucreler[c] if c < len(hucreler) else "")
                    if r == 0:
                        for run in hucre.paragraphs[0].runs:
                            run.bold = True
            continue
        if s.startswith("## "):
            doc.add_heading(s[3:].strip(), level=2)
        elif s.startswith("# "):
            doc.add_heading(s[2:].strip(), level=1)
        elif s.startswith(">"):
            p = doc.add_paragraph()
            _satir_ici(p, s.lstrip("> ").strip())
            for run in p.runs:
                run.italic = True
        elif re.match(r"\s*[-*] \[( |x|X)\] ", s):
            isaretli = re.match(r"\s*[-*] \[(x|X)\] ", s) is not None
            p = doc.add_paragraph()
            _satir_ici(p, f"{KUTU_DOLU if isaretli else KUTU_BOS} " + re.sub(r"\s*[-*] \[( |x|X)\] ", "", s, count=1))
        elif re.match(r"\s*[-*] ", s):
            p = doc.add_paragraph(style="List Bullet")
            _satir_ici(p, re.sub(r"\s*[-*] ", "", s, count=1))
        else:
            _satir_ici(doc.add_paragraph(), s.strip())
        i += 1


def belge_olustur(yanit: dict, bugun: date | None = None) -> bytes:
    """Kontrol listesi yanıtından .docx baytları."""
    bugun = bugun or date.today()
    t = yanit["tesvik"]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading(f"Başvuru dosyası: {t['baslik']}", level=0)
    bilgiler = [("Kurum", t.get("kurum")), ("Başvuru yeri", t.get("basvuru_yeri")),
                ("Başvuru dönemi", t.get("basvuru_suresi")), ("Resmi sayfa", t.get("kaynak_url")),
                ("Hazırlanma tarihi", bugun.strftime("%d.%m.%Y"))]
    for etiket, deger in bilgiler:
        if deger:
            p = doc.add_paragraph()
            p.add_run(f"{etiket}: ").bold = True
            p.add_run(str(deger))
    if t.get("aktif_mi") is False:
        doc.add_paragraph("Not: Bu program kayıtlarımıza göre şu an başvuruya kapalı görünüyor.").runs[0].bold = True

    doc.add_heading("Başvuru dönemleri", level=1)
    cagrilar = yanit.get("cagrilar") or []
    if not cagrilar:
        doc.add_paragraph("Doğrulanmış başvuru dönemi kaydı yok; tarihleri kurumun resmi sayfasından kontrol edin.")
    for c in cagrilar:
        p = doc.add_paragraph(style="List Bullet")
        tarih = " – ".join(x for x in (c.get("acilis"), c.get("kapanis")) if x) or "tarih duyurulmadı"
        p.add_run(f"{c['ad']}: ").bold = True
        p.add_run(f"{tarih} ({c.get('durum_metni', '')})")
        if c.get("notlar"):
            p.add_run(f". {c['notlar']}")
        p.add_run(f" Kaynak: {c['kaynak_url']} (doğrulama {c['dogrulama_tarihi']})")

    doc.add_heading(f"Kontrol listesi ({yanit.get('tamamlanan', 0)}/{yanit.get('toplam', 0)} tamamlandı)", level=1)
    maddeler = yanit.get("maddeler") or []
    if not maddeler:
        doc.add_paragraph(yanit.get("uyari") or "Bu program için şart/belge bilgisi yok.")
    for tur in ("sart", "belge", "basvuru"):
        grup = [m for m in maddeler if m["tur"] == tur]
        if not grup:
            continue
        doc.add_heading(TUR_BASLIK[tur], level=2)
        for m in grup:
            doc.add_paragraph(f"{KUTU_DOLU if m.get('isaretli') else KUTU_BOS} {m['metin']}")

    doc.add_heading("Başvuru ön taslağı", level=1)
    if yanit.get("taslak"):
        markdown_ekle(doc, yanit["taslak"])
    else:
        doc.add_paragraph("Henüz taslak oluşturulmadı. Panelde 'Taslak oluştur' ile yazılan metin buraya eklenir.")

    alt = doc.add_paragraph()
    alt.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = alt.add_run(UYARI)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    tampon = io.BytesIO()
    doc.save(tampon)
    return tampon.getvalue()


def dosya_adi(baslik: str) -> tuple[str, str]:
    """(ASCII yedek ad, UTF-8 ad) — Content-Disposition için."""
    temel = re.sub(r"[^\w\s-]", "", baslik, flags=re.UNICODE).strip()[:60] or "basvuru"
    utf8 = re.sub(r"\s+", "_", temel) + ".docx"
    ascii_ = (utf8.translate(str.maketrans("çğıöşüÇĞİÖŞÜ", "cgiosuCGIOSU")).encode("ascii", "ignore").decode()
              or "basvuru.docx")
    return ascii_, utf8


def _self_test() -> int:
    ornek = {
        "tesvik": {"id": 1, "baslik": "Yurt Dışı Marka Tescil Desteği (5973 sayılı Karar m.4)", "kurum": "Ticaret Bakanlığı",
                   "kaynak_url": "https://ticaret.gov.tr/x", "basvuru_yeri": "İBGS / DYS", "basvuru_suresi": "6 ay",
                   "aktif_mi": True},
        "maddeler": [{"anahtar": "s:1", "tur": "sart", "metin": "Şirket olmak", "isaretli": True},
                     {"anahtar": "b:1", "tur": "belge", "metin": "Fatura", "isaretli": False}],
        "tamamlanan": 1, "toplam": 2,
        "cagrilar": [{"ad": "2026-2", "acilis": "2026-09-01", "kapanis": "2026-10-31", "durum_metni": "Başvuruya açık",
                      "notlar": None, "kaynak_url": "https://k", "dogrulama_tarihi": "2026-10-08"}],
        "taslak": ("> **Ön taslak**\n\n## 1. İşletme tanıtımı\nMetin **kalın** [DOLDURUN: yıl]\n- madde\n- [ ] belge\n"
                   "- [x] tamam\n\n| Kalem | Tutar |\n|---|---|\n| Personel | [DOLDURUN] |\n"),
    }
    b = belge_olustur(ornek, bugun=date(2026, 10, 8))
    d = Document(io.BytesIO(b))
    metin = "\n".join(p.text for p in d.paragraphs)
    basliklar = [p.text for p in d.paragraphs if p.style.name.startswith(("Heading", "Title"))]
    kontroller = [
        ("geçerli docx (zip imzası)", b[:2] == b"PK"),
        ("başlıklar", basliklar[:2] == ["Başvuru dosyası: Yurt Dışı Marka Tescil Desteği (5973 sayılı Karar m.4)",
                                        "Başvuru dönemleri"] and "1. İşletme tanıtımı" in basliklar),
        ("kontrol listesi kutuları", f"{KUTU_DOLU} Şirket olmak" in metin and f"{KUTU_BOS} Fatura" in metin),
        ("taslak onay kutuları", f"{KUTU_BOS} belge" in metin and f"{KUTU_DOLU} tamam" in metin),
        ("tablo (ayırıcı satır atlandı)", len(d.tables) == 1 and len(d.tables[0].rows) == 2
         and d.tables[0].cell(1, 0).text == "Personel"),
        ("kalın işaretleri metne sızmadı", "**" not in metin),
        ("çağrı satırı ve kaynak", "2026-09-01 – 2026-10-31" in metin and "https://k" in metin),
        ("dosya adı: ASCII yedeği ve UTF-8 ad", dosya_adi("Yurt Dışı Marka: Tescil")
         == ("Yurt_Disi_Marka_Tescil.docx", "Yurt_Dışı_Marka_Tescil.docx")),
        ("taslaksız da üretilir", Document(io.BytesIO(belge_olustur({**ornek, "taslak": None}))).paragraphs[0].text
         .startswith("Başvuru dosyası")),
    ]
    for ad, ok in kontroller:
        print(("  ✓ " if ok else "  ✗ ") + ad)
    gecen = sum(ok for _, ok in kontroller)
    print(f"self-test: {gecen}/{len(kontroller)} geçti")
    return 0 if gecen == len(kontroller) else 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    if "--self-test" in sys.argv:
        sys.exit(_self_test())
    print("Kullanım: python -m app.basvuru_docx --self-test")
