"""Teşvik bulunduktan sonraki yardım katmanının (kontrol listesi, sihirbaz, şablon taslak, resmi form şablonu, Word)
ajanlarla uçtan uca denenmesi — 2026-10-09. Ücretli API çağrısı YOK.

  python docs/olcum/2026-10-09-taslak-ajan/taslak_ajan.py hazirla   -> paket_NN.md (kullanıcı ajanına), taslak_bos_NN.md
  (kullanıcı ajanı: cevap_NN.json + deneyim_NN.md yazar)
  python docs/olcum/2026-10-09-taslak-ajan/taslak_ajan.py uret      -> taslak_NN.md + otomatik denetim (denetim.json)
  (hakem ajanı: degerlendirme_NN.json yazar)
  python docs/olcum/2026-10-09-taslak-ajan/taslak_ajan.py ozet      -> ozet.md (otomatik denetim + hakem puanları)
  python docs/olcum/2026-10-09-taslak-ajan/taslak_ajan.py --self-test

Akış uygulamadakiyle aynı fonksiyonlardan geçer: _yanit (kontrol listesi), program_cagrilari, sablon_taslak.sorular /
uret, TaslakCevaplari (API'nin şema doğrulaması), basvuru_docx.belge_olustur (Word indirme).
"""
from __future__ import annotations

import io
import json
import os
import re
import sys
from datetime import date

KOK = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
sys.path.insert(0, KOK)
sys.path.insert(0, os.path.join(KOK, "docs", "olcum", "2026-10-08-taslak"))
OUT = os.path.dirname(os.path.abspath(__file__))

# (tesvik_id, profil, kullanıcının bildikleri — ajan bunlara dayanarak cevaplar; "Bilmediğin" kısmı bilinçli boşluk)
VAKALAR = [
    (34, dict(sektor="imalat", bolge="İzmir", calisan_sayisi=12, yillik_ciro=60e6, nace_kodu="13.20", sirket_turu="limited",
              hedefler=["arge", "ihracat"]),
     "Ege Örme Tekstil Ltd. Şti., İzmir, 12 çalışan, yıllık ciro yaklaşık 60 milyon TL. Örme kumaş üretiyorsunuz. Proje "
     "fikri: atık pamuk liflerinden geri dönüştürülmüş iplikle örme kumaş. Sorun: geri dönüştürülmüş lif oranı artınca "
     "kumaşın kopma mukavemeti düşüyor. Hedef: %40 geri dönüştürülmüş lif oranında kopma mukavemetinin standart kumaşın "
     "en az %90'ı olması. Ekip: 1 tekstil mühendisi (proje yürütücüsü olur), 1 teknisyen; Ege Üniversitesi'nden bir "
     "hocayla görüştünüz ama anlaşma yok. Süre 18 ay, Ocak 2027'de başlamak istiyorsunuz. Bütçe tahmini: personel "
     "1.800.000 TL, hammadde ve sarf malzeme 600.000 TL, dış laboratuvar test ve analiz 450.000 TL, numune örme makinesi "
     "900.000 TL. Avrupa'daki 3 mevcut müşteriniz (moda markaları) bu kumaşa ilgi gösterdi. Rakipleriniz İtalyan "
     "üreticiler. Bilmediğin: konudaki bilimsel literatürün ayrıntısı, patent alınabilir mi, şirketin daha önce TÜBİTAK "
     "projesi olup olmadığını muhasebe biliyor ama sen hatırlamıyorsun."),
    (44, dict(sektor="arge", bolge="İstanbul", calisan_sayisi=12, yillik_ciro=20e6, nace_kodu="62.01", sirket_turu="limited",
              trl=5, hedefler=["arge"]),
     "Bulut Muhasebe Yazılım Ltd. Şti., İstanbul, 12 çalışan, ciro 20 milyon TL. KOBİ'lere bulut muhasebe yazılımı "
     "satıyorsunuz (1.400 abone). İlk TÜBİTAK başvurunuz. Proje: gelen e-faturaları gider hesaplarına otomatik "
     "sınıflandıran makine öğrenmesi modülü. Bugün müşteriler bunu elle yapıyor; sizin kural tabanlı denemeniz %70 doğru. "
     "Hedef %95 doğruluk ve fatura başına 2 saniyenin altında işlem. 12 ay, Şubat 2027 başlangıç. Bütçe: personel "
     "(2 yazılımcı + 1 veri bilimci, kısmi) 1.400.000 TL, bulut/GPU ve yazılım lisansı 250.000 TL, üniversiteden "
     "danışmanlık 200.000 TL. Ticarileşme: mevcut abonelere ek modül olarak aylık ücretle. Bilmediğin: hangi ML yöntemi "
     "seçileceği henüz netleşmedi."),
    (8, dict(sektor="imalat", bolge="Konya", calisan_sayisi=28, yillik_ciro=85e6, nace_kodu="28.30", sirket_turu="limited",
             hedefler=["yatirim", "ihracat"]),
     "Konya Hassas Metal Ltd. Şti., 28 çalışan, ciro 85 milyon TL. Tarım makinelerine CNC ile yedek parça işliyorsunuz. "
     "Proje: 2 adet 5 eksen CNC işleme merkezi ve 1 koordinat ölçüm cihazı (CMM) alarak kapasiteyi %35 artırmak ve "
     "ihracatın satışlardaki payını %10'dan %25'e çıkarmak (müşteriler Almanya ve Romanya'da). Makineler 9.500.000 TL, "
     "CAM yazılımı 400.000 TL, operatör eğitimi 150.000 TL. 2 yeni CNC operatörü alınacak. Süre 12 ay, Mart 2027. "
     "Yurt içi 3 rakibiniz var (Konya ve Bursa'da). Bilmediğin: ürününüzün dış ticaret istatistikleri ve pazar büyüklüğü "
     "rakamları; yatırımın geri dönüş süresini hesaplamadınız."),
    (180, dict(sektor="hizmet", bolge="Van", calisan_sayisi=25, yillik_ciro=30e6, nace_kodu="55.10", sirket_turu="anonim",
               hedefler=["yatirim", "istihdam"]),
     "Van Gölü Turizm A.Ş., 25 çalışan, ciro 30 milyon TL, Van'da 40 odalı otel işletiyorsunuz. Proje: 30 odalı ek bina "
     "ile 70 odaya çıkmak. Yatırım tutarı 85.000.000 TL: inşaat 55.000.000, makine-teçhizat 18.000.000, mobilya-donanım "
     "12.000.000 TL. 15 yeni istihdam. Mart 2027 - Haziran 2028. Arsa şirketin mülkü. Bilmediğin: teşvik belgesi "
     "başvurusunun teknik ayrıntıları, finansmanın ne kadarının kredi olacağı."),
    (174, dict(sektor="arge", bolge="Ankara", calisan_sayisi=0, yillik_ciro=0, nace_kodu="62.01", sirket_turu="yok", trl=4,
               hedefler=["arge"]),
     "Ankara'da 26 yaşında bir yazılım mühendisisin; şirket kurmadın, bir arkadaşınla (ortak kurucu) birlikte "
     "çalışıyorsun. Fikir: işitme engelliler için Türk İşaret Dili'ni kameradan metne çeviren mobil uygulama. Laboratuvar "
     "ortamında 300 işaretle %80 tanıma yapan bir prototipin var (TRL 4). 12 ayda 1.000 işarete çıkıp 50 kullanıcılı "
     "pilot yapmak istiyorsun. Harcama planı: yazılım geliştirme hizmeti 300.000 TL, test cihazları 120.000 TL, pazar "
     "doğrulama ve kullanıcı testleri 80.000 TL. Bilmediğin: şirketi ne zaman kuracağın, ticari model (abonelik mi kurum "
     "satışı mı)."),
    (1, dict(sektor="hizmet", bolge="Eskişehir", calisan_sayisi=0, yillik_ciro=0, nace_kodu="10.71", sirket_turu="yok",
             hedefler=["finansman", "istihdam"]),
     "Eskişehir'de 34 yaşında bir kadınsın; Uygulamalı Girişimcilik Eğitimi sertifikan var, henüz işletme kurmadın. "
     "Butik pastane açacaksın (NACE 10.71). Kuruluş planı: şahıs işletmesi olarak Ocak 2027'de kuruluş, 2 çalışan. "
     "Giderler: fırın ve soğutma makineleri 350.000 TL, dükkân dekorasyonu 150.000 TL, ilk 6 ay kira 180.000 TL. İlk "
     "yıl satış hedefi 2.400.000 TL. Bilmediğin: hangi giderlerin desteklendiği, iş planının nasıl yazılacağı."),
    (162, dict(sektor="e-ticaret", bolge="Trabzon", calisan_sayisi=3, yillik_ciro=4e6, nace_kodu="47.91",
               sirket_turu="limited", hedefler=["ihracat"]),
     "Karadeniz Fındık Gıda Ltd. Şti., Trabzon, 3 çalışan, ciro 4 milyon TL. Fındık ezmesi ve kavrulmuş fındığı Amazon "
     "Almanya ve Etsy'de satıyorsunuz; geçen yıl yurt dışı satış 900.000 TL. 2027 planı: pazaryeri reklamı 240.000 TL, "
     "pazaryeri komisyonları 120.000 TL, Almanya'da sipariş karşılama (fulfillment) deposu 60.000 TL. Hedef 2027 yurt dışı "
     "satış 2.500.000 TL. İhracatçı birliği üyesisiniz. Bilmediğin: E-İhracat Destekleri'nde başvurunun hangi sırayla "
     "yapılacağı."),
    (2, dict(sektor="arge", bolge="İstanbul", calisan_sayisi=12, yillik_ciro=20e6, nace_kodu="62.01", sirket_turu="limited",
             trl=8, hedefler=["finansman", "arge"]),
     "Sohbet Yazılım Ltd. Şti., İstanbul, 12 çalışan, ciro 20 milyon TL. Yapay zekâ destekli müşteri hizmetleri asistanı "
     "ürününüz 40 kurumsal müşteride çalışıyor (TRL 8). Krediyle: GPU sunucu 1.800.000 TL, yapay zekâ yazılım lisansları "
     "400.000 TL, veri etiketleme hizmeti 300.000 TL. Teknogirişim Rozeti'niz var mı bilmiyorsun. Kesin teminat mektubu "
     "için bankayla görüşmediniz."),
    (185, dict(sektor="hizmet", bolge="Hatay", calisan_sayisi=4, yillik_ciro=3e6, nace_kodu="56.10", sirket_turu="sahis",
               hedefler=["istihdam"]),
     "Hatay'da 4 çalışanlı lokanta (şahıs işletmesi). Kasım 2026'da 3 kişi işe alacaksın: 22 ve 25 yaşında iki genç erkek "
     "ve 35 yaşında bir kadın; üçü de son 6 ayda kayıtlı çalışmamış, ikisi İŞKUR'a kayıtlı. Aylık brüt ücret asgari ücret. "
     "Bilmediğin: teşvikin hangi ayda başlayacağı ve sana kaç TL kazandıracağı."),
    (160, dict(sektor="tarim", bolge="Konya", calisan_sayisi=3, yillik_ciro=2.5e6, arazi_buyuklugu_dekar=180,
               urun_turu="buğday", tarim_kategori="tahil_baklagil", sirket_turu="sahis", hedefler=["makine"]),
     "Konya'da 180 dekar buğday eken bir çiftçisin (şahıs). ÇKS kaydın var. Bu sezon sertifikalı tohum kullanacaksın. "
     "Traktörün eski; yeni mibzer almak istiyorsun (180.000 TL). Bilmediğin: hububat desteğinin ne zaman ve nasıl "
     "ödendiği, mibzerin bu destekle ilgisi."),
    (182, dict(sektor="arge", bolge="İzmir", calisan_sayisi=18, yillik_ciro=35e6, nace_kodu="62.01", sirket_turu="limited",
               hedefler=["ihracat"]),
     "Ege Yazılım Ltd. Şti., İzmir, 18 çalışan, ciro 35 milyon TL, gelirin %30'u ABD'deki SaaS abonelerinden. 2027'de "
     "ABD'de büyümek için: yurt dışı dijital reklam 600.000 TL, Web Summit fuarı katılımı 350.000 TL, ABD pazar araştırması "
     "raporu 150.000 TL. Bilmediğin: hizmet ihracatçısı olarak hangi kayıtların gerektiği."),
    (3, dict(sektor="imalat", bolge="Kayseri", calisan_sayisi=22, yillik_ciro=40e6, nace_kodu="31.09", sirket_turu="limited",
             hedefler=["verimlilik"]),
     "Kayseri Ahşap Mobilya Ltd. Şti., 22 çalışan, ciro 40 milyon TL. Siparişleri Excel'de takip ediyorsunuz, üretimde "
     "gecikmeler var. Proje: ERP + üretim takip (MES) yazılımı ve atölyeye 6 tablet/barkod okuyucu. Yazılım 700.000 TL, "
     "donanım 300.000 TL. Hedef: teslim süresini 21 günden 14 güne indirmek. Bilmediğin: dijital dönüşüm danışmanı "
     "raporu gerekip gerekmediği."),
]

CEVAP_SEMASI = {
    "proje_adi": "metin (≤200)", "proje_ozeti": "metin (≤2000)", "gerekce": {"<soru anahtarı>": "metin (≤2000)"},
    "faaliyetler": [{"ad": "metin", "baslangic": "YYYY-AA veya null", "bitis": "YYYY-AA veya null"}],
    "butce": [{"kalem": "metin", "tutar": "sayı (TL)"}], "destek_orani": "oran_secenekleri'nden biri veya null",
    "ciktilar": {"<soru anahtarı veya 'olcum'>": "metin"},
}


def _hazir(n: int):
    from app.models import FinancialProfile, SessionLocal, Tesvik
    from app.rag import profil_sozlugu
    tid, prof, bilgi = VAKALAR[n - 1]
    t = SessionLocal().get(Tesvik, tid)
    p = FinancialProfile(**prof)
    return t, p, profil_sozlugu(p), bilgi


def hazirla():
    from app import sablon_taslak
    from app.basvuru_listesi import _yanit
    from app.cagrilar import program_cagrilari
    for n in range(1, len(VAKALAR) + 1):
        t, _, profil, bilgi = _hazir(n)
        y = _yanit(t, None)
        s = sablon_taslak.sorular(t)
        s.pop("cevaplar", None)
        bos = sablon_taslak.uret(t, profil, y["maddeler"], program_cagrilari(t))
        io.open(os.path.join(OUT, f"taslak_bos_{n:02}.md"), "w", encoding="utf-8").write(bos)
        madde = "\n".join(f"- [{m['tur']}{', kural' if m.get('kural') else ''}] {m['metin']}" for m in y["maddeler"])
        cag = "\n".join(f"- {c['ad']}: {c['durum_metni']}" for c in y["cagrilar"]) or "- (kayıtlı dönem yok)"
        paket = (f"# Vaka {n:02}: {t.baslik} — {t.kurum}\n\n## Senin işletmen (yalnız bunları biliyorsun)\n{bilgi}\n\n"
                 f"## Uygulamanın gösterdiği program bilgisi\nBaşvuru yeri: {t.basvuru_yeri or '-'}\nDönem: {t.basvuru_suresi or '-'}\n"
                 f"Tutar/oran: {t.tesvil_tutari or '-'}\nKaynak: {t.kaynak_url}\n\nKontrol listesi:\n{madde}\n\nBaşvuru dönemleri:\n{cag}\n\n"
                 f"## Sihirbaz soruları (GET /api/basvuru-listesi/{t.id}/taslak-sorulari yanıtı)\n```json\n"
                 f"{json.dumps(s, ensure_ascii=False, indent=1, default=str)}\n```\n\n## Cevap biçimi\n```json\n"
                 f"{json.dumps(CEVAP_SEMASI, ensure_ascii=False, indent=1)}\n```\n")
        io.open(os.path.join(OUT, f"paket_{n:02}.md"), "w", encoding="utf-8").write(paket)
        print(f"[{n:2}] {t.id:3} {t.baslik[:55]:55} form={'evet' if s['form'] else '-':4} sorular={len(s['gerekce'])}+{len(s['cikti'])} "
              f"oran={s['oran_secenekleri']} limit={s['ust_limit']}")


# ----------------------------------------------------------------------------------------- otomatik denetim
def _metinler(c: dict) -> list[str]:
    m = [c.get("proje_adi", ""), *c.get("gerekce", {}).values(), *c.get("ciktilar", {}).values()]
    m += [f["ad"] for f in c.get("faaliyetler", [])] + [b["kalem"] for b in c.get("butce", [])]
    return [x.strip() for x in m if x and x.strip()]


def denetle(n: int, cevap: dict) -> tuple[dict, str]:
    import taslak_olcum as TO
    from app import sablon_taslak
    from app.basvuru_docx import belge_olustur
    from app.basvuru_listesi import TaslakCevaplari, _yanit
    from app.basvuru_taslagi import baglam
    from app.cagrilar import program_cagrilari
    from app.form_sablonlari import form as resmi_form
    t, _, profil, _ = _hazir(n)
    k: dict[str, object] = {}
    try:
        c = TaslakCevaplari.model_validate(cevap).model_dump()
        k["sema_gecerli"] = True
    except Exception as e:  # API 422 döner; ajan cevabı uygulamada da reddedilirdi
        return {"sema_gecerli": False, "sema_hatasi": str(e)[:400]}, ""
    y = _yanit(t, None)
    cag = program_cagrilari(t)
    metin = sablon_taslak.uret(t, profil, y["maddeler"], cag, c)
    s = sablon_taslak.sorular(t)
    eksik = [x for x in _metinler(c) if x not in metin]
    k["cevaplar_taslakta"] = not eksik
    k["taslakta_olmayan_cevaplar"] = [x[:80] for x in eksik]
    top = sum(b["tutar"] for b in c["butce"])
    beklenen = None
    if c["destek_orani"] and top:
        beklenen = min(top * c["destek_orani"] / 100, s["ust_limit"]) if s["ust_limit"] else top * c["destek_orani"] / 100
    k["butce_toplami"] = top
    k["butce_hesabi_dogru"] = (not top or sablon_taslak._tl(top) in metin) and (beklenen is None or sablon_taslak._tl(beklenen) in metin)
    k["beklenen_destek"] = beklenen
    kaynak = (baglam(t, profil) + json.dumps(cag, ensure_ascii=False, default=str) + json.dumps(resmi_form(t.id) or {}, ensure_ascii=False)
              + json.dumps(c, ensure_ascii=False) + " ".join(f"{k_}: {v}" for k_, v in profil.items())
              # kodun hesapladığı toplam ve destek (doğruluğu butce_hesabi_dogru ile ayrıca denetlenir)
              + f" {sablon_taslak._tl(top)} " + (f"{sablon_taslak._tl(beklenen)} {sablon_taslak._tl(top * c['destek_orani'] / 100)}"
                                                if beklenen is not None else ""))
    govde = re.sub(r"^## 6\..*", "", metin, flags=re.S | re.M)
    k["uydurma_sayilar"] = TO.denetle(govde, kaynak)["baglamda_olmayan_sayilar"]
    f = resmi_form(t.id)
    if f:
        sira = [b[1] for b in f["bolumler"]]
        bulunan = [b for b in sira if f"### {b}" in metin]
        k["form_basliklari_sirali"] = bulunan == sira and all(metin.index(f"### {a}") < metin.index(f"### {b}") for a, b in zip(sira, sira[1:]))
    kural = [m["metin"] for m in y["maddeler"] if m.get("kural")]
    k["kurallar_uyari_olarak"] = all(f"- ⚠ {x}" in metin and f"- [ ] {x}" not in metin for x in kural)
    ters = [f["ad"] for f in c["faaliyetler"] if f["baslangic"] and f["bitis"] and f["bitis"] < f["baslangic"]]
    k["ters_tarihli_faaliyet"] = ters
    k["doldurun_sayisi"] = metin.count("[DOLDURUN")
    yy = dict(y)
    yy.update(taslak=metin, taslak_tarihi="2026-10-09T12:00:00")
    try:
        from docx import Document
        d = Document(io.BytesIO(belge_olustur(yy, bugun=date(2026, 10, 9))))
        govde_docx = "\n".join(p.text for p in d.paragraphs) + "\n".join(c_.text for tb in d.tables for r in tb.rows for c_ in r.cells)
        k["word_uretildi"] = True
        k["word_proje_adi_iceriyor"] = (c["proje_adi"] or "") in govde_docx
    except Exception as e:
        k["word_uretildi"] = False
        k["word_hatasi"] = f"{type(e).__name__}: {e}"[:300]
    return k, metin


def uret():
    sonuc = {}
    for n in range(1, len(VAKALAR) + 1):
        yol = os.path.join(OUT, f"cevap_{n:02}.json")
        if not os.path.exists(yol):
            print(f"[{n:2}] cevap yok, atlandı")
            continue
        k, metin = denetle(n, json.load(io.open(yol, encoding="utf-8")))
        if metin:
            io.open(os.path.join(OUT, f"taslak_{n:02}.md"), "w", encoding="utf-8").write(metin)
        sonuc[n] = k
        print(f"[{n:2}] " + ", ".join(f"{a}={b}" for a, b in k.items() if a not in ("taslakta_olmayan_cevaplar",)))
    json.dump(sonuc, io.open(os.path.join(OUT, "denetim.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1, default=str)


PUANLAR = ["forma_aktarilabilirlik", "programa_uygunluk", "kullanici_bilgisinin_yansimasi", "ic_tutarlilik",
           "eksiklerin_isaretlenmesi", "dogruluk_yaniltmama"]


def ozet():
    d = json.load(io.open(os.path.join(OUT, "denetim.json"), encoding="utf-8"))
    sat = ["| # | program | şema | cevaplar taslakta | bütçe hesabı | uydurma sayı | form başlıkları | kural uyarısı | Word | DOLDURUN | hakem ort. (1-5) | aktarılabilir |",
           "|---|---|---|---|---|---|---|---|---|---|---|---|"]
    tum = {p: [] for p in PUANLAR}
    for n in range(1, len(VAKALAR) + 1):
        k = d.get(str(n))
        if not k:
            continue
        h_yol = os.path.join(OUT, f"degerlendirme_{n:02}.json")
        h = json.load(io.open(h_yol, encoding="utf-8")) if os.path.exists(h_yol) else {}
        p = h.get("puanlar", {})
        for a in PUANLAR:
            if isinstance(p.get(a), (int, float)):
                tum[a].append(p[a])
        ort = round(sum(p[a] for a in PUANLAR if a in p) / max(1, sum(a in p for a in PUANLAR)), 1) if p else "-"
        e = lambda x: "✓" if x else ("-" if x is None else "✗")  # noqa: E731
        sat.append(f"| {n} | {VAKALAR[n-1][0]} | {e(k.get('sema_gecerli'))} | {e(k.get('cevaplar_taslakta'))} | "
                   f"{e(k.get('butce_hesabi_dogru'))} | {', '.join(map(str, k.get('uydurma_sayilar', []))) or '-'} | "
                   f"{e(k.get('form_basliklari_sirali'))} | {e(k.get('kurallar_uyari_olarak'))} | {e(k.get('word_uretildi'))} | "
                   f"{k.get('doldurun_sayisi', '-')} | {ort} | {h.get('resmi_forma_aktarilabilir', '-')} |")
    sat += ["", "Hakem puan ortalamaları: " + ", ".join(f"{a} {sum(v)/len(v):.1f}" for a, v in tum.items() if v)]
    io.open(os.path.join(OUT, "ozet.md"), "w", encoding="utf-8").write("\n".join(sat) + "\n")
    print("\n".join(sat))


def _self_test() -> int:
    from app import sablon_taslak  # noqa: F401
    k = [("12 vaka, tekil program", len(VAKALAR) == 12 and len({v[0] for v in VAKALAR}) == 12),
         ("cevap metinleri toplanır", _metinler({"proje_adi": "P", "gerekce": {"a": " x "}, "faaliyetler": [{"ad": "F"}],
                                                  "butce": [{"kalem": "K", "tutar": 1}], "ciktilar": {}}) == ["P", "x", "F", "K"])]
    t, _, _, _ = _hazir(1)
    cev = {"proje_adi": "Deneme projesi", "gerekce": {"B2": "Mevcut durum metni"},
           "faaliyetler": [{"ad": "Prototip", "baslangic": "2027-01", "bitis": "2027-06"}],
           "butce": [{"kalem": "personel", "tutar": 1_000_000}], "destek_orani": None, "ciktilar": {}}
    d, metin = denetle(1, cev)
    k += [("şema + cevaplar + bütçe", d["sema_gecerli"] and d["cevaplar_taslakta"] and d["butce_hesabi_dogru"]),
          ("form başlıkları sıralı (1501)", d.get("form_basliklari_sirali") is True),
          ("Word üretildi", d["word_uretildi"] and d["word_proje_adi_iceriyor"]),
          ("geçersiz şema yakalanır", denetle(1, {"destek_orani": 150})[0]["sema_gecerli"] is False),
          ("ters tarih yakalanır", denetle(1, {**cev, "faaliyetler": [{"ad": "X", "baslangic": "2027-05", "bitis": "2027-01"}]})[0]
           ["ters_tarihli_faaliyet"] == ["X"])]
    for ad, ok in k:
        print(f"  {'OK ' if ok else 'HATA'} {ad}")
    g = sum(ok for _, ok in k)
    print(f"\nself-test: {g}/{len(k)} geçti")
    return 0 if g == len(k) else 1


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    a = sys.argv[1:]
    if "--self-test" in a:
        sys.exit(_self_test())
    {"hazirla": hazirla, "uret": uret, "ozet": ozet}.get(a[0] if a else "", lambda: print(__doc__))()
