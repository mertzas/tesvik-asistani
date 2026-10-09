"""Başvuru dosyası Word çıktısı (app/basvuru_docx.py, GET /api/basvuru-listesi/{id}/docx), 2026-10-08."""
import io
import re
import subprocess
import sys
from datetime import date, timedelta
from pathlib import Path
from urllib.parse import unquote

from docx import Document

from app.models import BasvuruTakibi, Organization, Tesvik, TesvikCagrisi

KOK = Path(__file__).resolve().parent.parent


def _hazirla(client, db_session, test_user_token):
    db_session.add(Tesvik(id=77, kurum="Ticaret Bakanlığı", baslik="Yurt Dışı Marka Tescil Desteği (5973 m.4)",
                          ozet="o", detay="d", aktif_mi=True, kaynak_url="https://t/77",
                          basvuru_sartlari=["Şirket olmak"], gerekli_belgeler=["Fatura", "Yurt içi marka tescil belgesi"],
                          basvuru_yeri="İBGS / DYS", basvuru_suresi="Ödeme belgesinden itibaren 6 ay"))
    db_session.add(TesvikCagrisi(tesvik_id=77, ad="Sürekli", acilis=date.today() - timedelta(days=3),
                                 kapanis=date.today() + timedelta(days=30), kaynak_url="https://k/77",
                                 dogrulama_tarihi=date.today()))
    db_session.commit()
    return {"Authorization": f"Bearer {test_user_token}"}


def test_word_dosyasi_liste_cagri_ve_taslak_icerir(client, db_session, test_user_token):
    h = _hazirla(client, db_session, test_user_token)
    liste = client.get("/api/basvuru-listesi/77", headers=h).json()
    ilk = liste["maddeler"][0]["anahtar"]
    belge = next(m["anahtar"] for m in liste["maddeler"] if m["tur"] == "belge")
    client.put("/api/basvuru-listesi/77", headers=h, json={"isaretli": [belge], "uygunluk": {ilk: "evet"}})
    kayit = db_session.query(BasvuruTakibi).one()
    kayit.taslak = "## 1. İşletme tanıtımı\nMetin [DOLDURUN: yıl]\n\n| Kalem | Tutar |\n|---|---|\n| Personel | [DOLDURUN] |"
    db_session.commit()

    r = client.get("/api/basvuru-listesi/77/docx", headers=h)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml")
    cd = r.headers["content-disposition"]
    assert 'filename="Yurt_Disi_Marka_Tescil_Destegi_5973_m4.docx"' in cd
    assert unquote(re.search(r"filename\*=UTF-8''(\S+)", cd).group(1)) == "Yurt_Dışı_Marka_Tescil_Desteği_5973_m4.docx"
    assert r.headers["cache-control"] == "no-store"

    d = Document(io.BytesIO(r.content))
    metin = "\n".join(p.text for p in d.paragraphs)
    # şart cevabıyla, belge kutuyla (kutular karışmaz)
    assert "Şirket olmak — ✔ sağlıyor" in metin and "☑ Fatura" in metin and "☑ Şirket olmak" not in metin
    assert "Sürekli:" in metin and "Başvuruya açık" in metin and "https://k/77" in metin
    assert "1. İşletme tanıtımı" in [p.text for p in d.paragraphs] and d.tables[0].cell(1, 0).text == "Personel"
    assert "resmi başvuru formu değildir" in metin


def test_taslaksiz_ve_takipsiz_de_uretilir(client, db_session, test_user_token):
    h = _hazirla(client, db_session, test_user_token)
    r = client.get("/api/basvuru-listesi/77/docx", headers=h)
    assert r.status_code == 200 and "Henüz taslak oluşturulmadı" in "\n".join(
        p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    assert db_session.query(BasvuruTakibi).count() == 0, "indirme takip kaydı açmamalı"


def test_kimliksiz_ve_olmayan_program(client, db_session, test_user_token):
    h = _hazirla(client, db_session, test_user_token)
    assert client.get("/api/basvuru-listesi/77/docx").status_code in (401, 403)
    assert client.get("/api/basvuru-listesi/999/docx", headers=h).status_code == 404


def test_baska_kurulusun_isaretleri_sizmaz(client, db_session, test_user_token):
    h = _hazirla(client, db_session, test_user_token)
    baska = Organization(name="Başka", email="baska@example.com")
    db_session.add(baska)
    db_session.flush()
    from app.basvuru_listesi import maddeler
    t = db_session.get(Tesvik, 77)
    db_session.add(BasvuruTakibi(org_id=baska.id, tesvik_id=77, isaretli=[m["anahtar"] for m in maddeler(t)],
                                 taslak="GİZLİ TASLAK"))
    db_session.commit()
    metin = "\n".join(p.text for p in Document(io.BytesIO(
        client.get("/api/basvuru-listesi/77/docx", headers=h).content)).paragraphs)
    assert "GİZLİ TASLAK" not in metin and "☑" not in metin


def test_self_test_bayragi():
    r = subprocess.run([sys.executable, "-m", "app.basvuru_docx", "--self-test"], cwd=KOK, stdin=subprocess.DEVNULL,
                       capture_output=True, text=True, encoding="utf-8")
    m = re.search(r"self-test: (\d+)/(\d+) geçti", r.stdout)
    assert r.returncode == 0 and m and m.group(1) == m.group(2) and int(m.group(2)) >= 9
